#!/usr/bin/env python3
"""Compare the native plugin with the unchanged Python exporter, using fake data."""
import copy
import datetime
import json
import sys
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT.parent))
import hvmtools
from openpyxl import load_workbook

OUT = ROOT / "build" / "parity"
NOW = "2026-10-29T00:00:00Z"


class Frozen(datetime.datetime):
    @classmethod
    def now(cls, tz=None):
        value = cls(2026, 10, 29, tzinfo=datetime.timezone.utc)
        return value.astimezone(tz) if tz else value.replace(tzinfo=None)


def ref(value):
    return value.get("id") if isinstance(value, dict) else value


def scope(source, cluster):
    inv = copy.deepcopy(source)
    cluster = next(c for c in inv["clusters"] if c["id"] == cluster["id"])
    hosts = {ref(s) for s in cluster.get("servers", [])}
    servers = [
        s for s in inv["servers"]
        if s["id"] in hosts or ref(s.get("parentServer")) in hosts
        or ref(s.get("cluster")) == cluster["id"]
        or ref(s.get("serverGroup")) == cluster["id"]
    ]
    ids = {s["id"] for s in servers}
    instances = [
        i for i in inv["instances"]
        if any(ref(s) in ids for s in i.get("servers", []))
        or any(isinstance(c, dict) and ref(c.get("server")) in ids for c in i.get("containers", []))
    ]
    ds_ids = {
        v.get("datastoreId") or ref(v.get("datastore"))
        for s in servers for v in s.get("volumes", [])
    }
    datastores = [
        d for d in inv["datastores"]
        if d["id"] in ds_ids or cluster["name"] in d.get("_clusters", [])
        or ref(d.get("cluster")) == cluster["id"]
    ]
    for d in datastores:
        d["_clusters"] = [cluster["name"]] if (
            cluster["name"] in d.get("_clusters", [])
            or ref(d.get("cluster")) == cluster["id"]) else []
    net_ids = {ref(n.get("network")) for s in servers for n in s.get("interfaces", [])}
    cloud_ids = {ref(cluster.get("zone")), ref(cluster.get("cloud"))}
    group_ids = {ref(cluster.get("site")), ref(cluster.get("group"))}
    for s in servers:
        cloud_ids.update([ref(s.get("zone")), ref(s.get("cloud"))])
        group_ids.add(ref(s.get("site")))
    group_ids.update(ref(i.get("group")) for i in instances)
    instance_ids = {str(i["id"]) for i in instances}
    inv.update(
        clusters=[cluster], servers=servers, instances=instances, datastores=datastores,
        networks=[n for n in inv["networks"] if n["id"] in net_ids],
        clouds=[c for c in inv["clouds"] if c["id"] in cloud_ids],
        groups=[g for g in inv["groups"] if g["id"] in group_ids],
        snapshots={k: v for k, v in inv["snapshots"].items() if str(k) in instance_ids},
        whoami={},
    )
    return inv


def generate():
    OUT.mkdir(parents=True, exist_ok=True)
    inv = hvmtools.redact(json.loads((ROOT.parent / "examples/sample-inventory.json").read_text()))
    cases = [("all", inv)] + [
        (f"cluster-{c['id']}", scope(inv, c)) for c in inv["clusters"]
    ]
    manifest = {"now": NOW, "inventory": inv, "cases": []}
    with patch.object(hvmtools.dt, "datetime", Frozen):
        for name, inventory in cases:
            sheets = hvmtools.build_sheets(inventory)
            hvmtools.write_workbook(sheets, OUT / f"python-{name}.xlsx", inventory, raw=True)
            manifest["cases"].append({"name": name, "inventory": inventory, "sheets": sheets})
        manifest["sheets"] = manifest["cases"][0]["sheets"]
    (OUT / "oracle.json").write_text(json.dumps(manifest, default=lambda d: d.isoformat()))
    print(f"Generated Python oracle: {len(cases)} inventories, all 13 tabs plus raw sheets.")


def font_key(font):
    return font.name, font.sz, bool(font.bold), color_key(font.color) if font.color else None


def color_key(color):
    return color.rgb[-6:] if color.type == "rgb" and isinstance(color.rgb, str) else None


def compare():
    manifest = json.loads((OUT / "oracle.json").read_text())
    cells = 0
    for case in manifest["cases"]:
        name = case["name"]
        expected = load_workbook(OUT / f"python-{name}.xlsx")
        actual = load_workbook(OUT / "native" / f"{name}.xlsx")
        assert actual.sheetnames == expected.sheetnames, name
        for title in expected.sheetnames:
            left, right = expected[title], actual[title]
            assert left.max_row == right.max_row and left.max_column == right.max_column, (name, title, "dimensions")
            assert left.freeze_panes == right.freeze_panes, (name, title, "freeze")
            assert left.auto_filter.ref == right.auto_filter.ref, (name, title, "filter")
            for col in left.column_dimensions:
                assert left.column_dimensions[col].width == right.column_dimensions[col].width, (name, title, col, "width")
            for row in left:
                for a in row:
                    b = right[a.coordinate]
                    label = (name, title, a.coordinate)
                    if isinstance(a.value, (int, float)) and not isinstance(a.value, bool):
                        assert isinstance(b.value, (int, float)) and abs(a.value - b.value) < 1e-8, (label, a.value, b.value)
                    else:
                        assert a.value == b.value, (label, a.value, b.value)
                    assert a.data_type == b.data_type, (label, "cell type", a.data_type, b.data_type)
                    assert a.number_format == b.number_format, (label, "number format")
                    assert font_key(a.font) == font_key(b.font), (label, "font")
                    assert a.fill.patternType == b.fill.patternType, (label, "fill")
                    assert color_key(a.fill.fgColor) == color_key(b.fill.fgColor), (label, "fill color")
                    assert a.alignment.vertical == b.alignment.vertical, (label, "alignment")
                    cells += 1
        actual.close()
        expected.close()
    print(f"Workbook parity verified: {len(manifest['cases'])} inventories, {cells:,} cells; values, types, styles, widths, filters and frozen panes match.")


if __name__ == "__main__":
    {"generate": generate, "compare": compare}[sys.argv[1]]()
